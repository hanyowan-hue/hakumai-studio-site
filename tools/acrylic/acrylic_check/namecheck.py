"""SVG に入っている参列者名と、Excel/CSV の名簿を突き合わせる。

前提：名前がテキストレイヤーとして残っていること。
画像に焼き込んだ（ラスタライズした / アウトライン化した）データからは
文字を読み取れないため、その場合は検査できない旨を報告する。
"""

from __future__ import annotations

import csv
import unicodedata
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

from .svgdoc import SVG_NS, INKSCAPE_NS, apply_matrix, identity, parse_transform, strip_ns


@dataclass
class SvgName:
    text: str
    normalized: str
    element_id: str
    position_mm: tuple[float, float] | None
    label_path: tuple[str, ...] = ()


@dataclass
class NameReport:
    svg_names: list[SvgName]
    roster: list[str]
    missing: list[str] = field(default_factory=list)      # 名簿にあるのに SVG にない
    extra: list[str] = field(default_factory=list)        # SVG にあるのに名簿にない
    duplicated_in_svg: list[tuple[str, int]] = field(default_factory=list)
    duplicated_in_roster: list[tuple[str, int]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not (self.missing or self.extra or self.duplicated_in_svg)


def normalize(text: str, strip_spaces: bool = False) -> str:
    """全角/半角・記号ゆれ・空白ゆれを吸収する。"""
    t = unicodedata.normalize("NFKC", text)
    t = t.replace("　", " ")
    t = " ".join(t.split())
    if strip_spaces:
        t = t.replace(" ", "")
    return t.strip()


def extract_names(
    svg_path: str,
    to_mm: np.ndarray | None = None,
    strip_spaces: bool = False,
) -> tuple[list[SvgName], list[str]]:
    """SVG の <text> から名前を取り出す。"""
    warnings: list[str] = []
    tree = ET.parse(svg_path)
    root = tree.getroot()
    names: list[SvgName] = []

    def walk(elem: ET.Element, ctm: np.ndarray, labels: tuple[str, ...]):
        tag = strip_ns(elem.tag)
        ctm = ctm @ parse_transform(elem.attrib.get("transform"))
        if tag == "g":
            name = elem.attrib.get(f"{{{INKSCAPE_NS}}}label") or elem.attrib.get("id") or ""
            if name:
                labels = (*labels, name)
        if tag == "text":
            raw = "".join(elem.itertext())
            text = normalize(raw, strip_spaces)
            if text:
                pos = None
                try:
                    x = float(elem.attrib.get("x", "nan"))
                    y = float(elem.attrib.get("y", "nan"))
                    if x == x and y == y:  # NaN でない
                        p = apply_matrix(np.array([[x, y]]), ctm)[0]
                        if to_mm is not None:
                            p = apply_matrix(np.array([p]), to_mm)[0]
                        pos = (float(p[0]), float(p[1]))
                except ValueError:
                    pos = None
                names.append(
                    SvgName(raw.strip(), text, elem.attrib.get("id", ""), pos, labels)
                )
            return  # text の中の tspan は itertext で回収済み
        for child in list(elem):
            walk(child, ctm, labels)

    walk(root, identity(), ())

    if not names:
        warnings.append(
            "SVG にテキスト要素が 1 つもありません。"
            "名前が画像に焼き込まれている / アウトライン化されていると照合できません。"
            "照合したい場合はテキストレイヤーを残したデータで書き出してください。"
        )
    return names, warnings


def load_roster(path: str, column: str | int | None = None, strip_spaces: bool = False) -> list[str]:
    """Excel(.xlsx) または CSV から名簿を読む。"""
    p = Path(path)
    suffix = p.suffix.lower()
    rows: list[list[str]]

    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("Excel を読むには openpyxl が必要です: pip install openpyxl") from exc
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.active
        rows = [
            ["" if v is None else str(v) for v in row]
            for row in ws.iter_rows(values_only=True)
        ]
    else:
        with p.open(encoding="utf-8-sig", newline="") as f:
            rows = [list(r) for r in csv.reader(f)]

    if not rows:
        return []

    col_index = 0
    header = rows[0]
    if isinstance(column, int):
        col_index = column
        body = rows
    elif isinstance(column, str):
        norm_header = [normalize(h) for h in header]
        if normalize(column) not in norm_header:
            raise ValueError(f"列 {column!r} が見つかりません。ヘッダー: {header}")
        col_index = norm_header.index(normalize(column))
        body = rows[1:]
    else:
        body = rows

    out: list[str] = []
    for row in body:
        if col_index >= len(row):
            continue
        value = normalize(str(row[col_index]), strip_spaces)
        if value:
            out.append(value)
    return out


def compare(svg_names: list[SvgName], roster: list[str]) -> NameReport:
    svg_counter = Counter(n.normalized for n in svg_names)
    roster_counter = Counter(roster)

    missing = sorted((roster_counter - svg_counter).elements())
    extra = sorted((svg_counter - roster_counter).elements())
    dup_svg = sorted((k, v) for k, v in svg_counter.items() if v > 1)
    dup_roster = sorted((k, v) for k, v in roster_counter.items() if v > 1)

    return NameReport(
        svg_names=svg_names,
        roster=roster,
        missing=missing,
        extra=extra,
        duplicated_in_svg=dup_svg,
        duplicated_in_roster=dup_roster,
    )
